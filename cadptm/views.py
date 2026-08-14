import logging
import uuid
from datetime import datetime
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST

from .forms import PatrimonioForm
from .models import (
    Auditoria,
    HistoricoAlteracaoPatrimonio,
    Patrimonio,
    SequenciaPatrimonio,
)

logger = logging.getLogger(__name__)

try:
    import pandas as pd
except ImportError:  # pragma: no cover
    pd = None

try:
    import qrcode  # noqa: F401
except ImportError:  # pragma: no cover
    qrcode = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Image,
        Paragraph,
        SimpleDocTemplate,
        Table,
        TableStyle,
    )
except ImportError:  # pragma: no cover
    A4 = None
    colors = None
    getSampleStyleSheet = None
    Image = None
    Paragraph = None
    SimpleDocTemplate = None
    Table = None
    TableStyle = None

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Border, Font, Side
except ImportError:  # pragma: no cover
    Workbook = None
    XLImage = None
    Font = None
    Border = None
    Side = None

try:
    import barcode
except ImportError:  # pragma: no cover
    barcode = None


def _require_export_dependencies():
    missing = []
    if Workbook is None or Font is None or XLImage is None:
        missing.append("openpyxl")
    if SimpleDocTemplate is None or Table is None or Paragraph is None or A4 is None:
        missing.append("reportlab")
    if pd is None:
        missing.append("pandas")
    return missing


@login_required
def dashboard(request):

    dados = Patrimonio.objects.values("status").annotate(total=Count("id"))

    resultado = {"ativo": 0, "baixado": 0}

    for item in dados:
        resultado[item["status"]] = item["total"]

    return JsonResponse(resultado)


@login_required
def etiqueta(request, id):
    item = Patrimonio.objects.filter(id=id).first()
    if not item:
        return HttpResponse(
            "Item não encontrado.", status=404, content_type="text/plain"
        )
    return render(request, "etiqueta.html", {"item": item})


@login_required
def exportar_excel_etiqueta(request, ids):
    ids_list = ids.split(",")
    itens = Patrimonio.objects.filter(id__in=ids_list)

    missing = _require_export_dependencies()
    if missing:
        logger.warning("Dependências de exportação ausentes: %s", ", ".join(missing))
        return HttpResponse(
            "Exportação indisponível no momento.",
            status=503,
            content_type="text/plain",
        )

    wb = Workbook()
    ws = wb.active

    linha = 1

    for item in itens:

        ws[f"A{linha}"] = "PATRIMÔNIO"
        ws[f"A{linha}"].font = Font(bold=True)

        ws[f"A{linha+1}"] = f"Controle: {item.controle}"
        ws[f"A{linha+2}"] = f"Patrimônio: {item.patrimonio}"
        ws[f"A{linha+3}"] = f"Material: {item.material}"
        ws[f"A{linha+4}"] = f"Local: {item.localizacao}"
        ws[f"A{linha+5}"] = f"Status: {item.status}"

        if item.qr_code:
            try:
                img = XLImage(item.qr_code.path)
                img.width = 100
                img.height = 100
                ws.add_image(img, f"E{linha+1}")
            except:
                pass

        linha += 8

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="etiquetas.xlsx"'
    wb.save(response)
    return response


@login_required
def exportar(request, formato, ids):

    ids_list = ids.split(",")
    itens = Patrimonio.objects.filter(id__in=ids_list)

    missing = _require_export_dependencies()
    if missing:
        logger.warning("Dependências de exportação ausentes: %s", ", ".join(missing))
        return HttpResponse(
            "Exportação indisponível no momento.",
            status=503,
            content_type="text/plain",
        )

    wb = Workbook()
    ws = wb.active

    # =========================
    # ✅ EXCEL - TABELA
    # =========================
    if formato == "excel":

        ws["A1"] = "RELATÓRIO DE PATRIMÔNIO"
        ws["A1"].font = Font(size=14, bold=True)

        headers = [
            "Controle",
            "Patrimônio",
            "Tipo",
            "Material",
            "Localização",
            "Status",
            "QR",
        ]
        ws.append(headers)

        linha = 3

        for item in itens:

            ws.cell(row=linha, column=1, value=item.controle)
            ws.cell(row=linha, column=2, value=item.patrimonio)
            ws.cell(row=linha, column=3, value=item.tipo)
            ws.cell(row=linha, column=4, value=item.material)
            ws.cell(row=linha, column=5, value=item.localizacao)
            ws.cell(row=linha, column=6, value=item.status)

            if item.qr_code:
                try:
                    img = XLImage(item.qr_code.path)
                    img.width = 60
                    img.height = 60
                    ws.row_dimensions[linha].height = 50
                    ws.add_image(img, f"G{linha}")
                except:
                    pass

            linha += 1

    # =========================
    # ✅ EXCEL ETIQUETA
    # =========================
    elif formato == "excel_etiqueta":

        linha = 1

        for item in itens:

            ws[f"A{linha}"] = "PATRIMÔNIO"
            ws[f"A{linha}"].font = Font(bold=True)

            ws[f"A{linha+1}"] = f"Controle: {item.controle}"
            ws[f"A{linha+2}"] = f"Patrimônio: {item.patrimonio}"
            ws[f"A{linha+3}"] = f"Material: {item.material}"
            ws[f"A{linha+4}"] = f"Local: {item.localizacao}"
            ws[f"A{linha+5}"] = f"Status: {item.status}"

            if item.qr_code:
                try:
                    img = XLImage(item.qr_code.path)
                    img.width = 100
                    img.height = 100
                    ws.add_image(img, f"E{linha+1}")
                except:
                    pass

            linha += 8

    # =========================
    # ✅ PDF
    # =========================
    elif formato == "pdf":

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()

        elementos = []

        # ✅ LOGO (coloque em /media/logo.png)
        try:
            logo = Image("media/logo.png", width=120, height=50)
            elementos.append(logo)
        except:
            pass

        # ✅ TÍTULO
        elementos.append(Paragraph("RELATÓRIO DE PATRIMÔNIO", styles["Title"]))

        # ✅ DATA E HORA
        agora = datetime.now().strftime("%d/%m/%Y %H:%M")
        elementos.append(Paragraph(f"Emitido em: {agora}", styles["Normal"]))

        dados = [["Controle", "Patrimônio", "Material", "Local", "Status", "QR"]]

        for item in itens:

            qr_path = item.qr_code.path if item.qr_code else ""

            try:
                qr_img = Image(qr_path, width=40, height=40)
            except:
                qr_img = ""

            dados.append(
                [
                    item.controle,
                    item.patrimonio,
                    item.material,
                    item.localizacao,
                    item.status,
                    qr_img,
                ]
            )

        tabela = Table(dados)

        tabela.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ]
            )
        )

        elementos.append(tabela)

        doc.build(elementos)
        buffer.seek(0)

        return HttpResponse(buffer, content_type="application/pdf")

    # =========================
    # ✅ RESPONSE FINAL (EXCEL)
    # =========================
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = f'attachment; filename="{formato}.xlsx"'

    output = BytesIO()
    wb.save(output)
    response.write(output.getvalue())

    return response


def _normalizar_valor_historico(valor):
    if valor is None:
        return ""
    if hasattr(valor, "isoformat"):
        return valor.isoformat()
    return str(valor)


def _registrar_historico_alteracao(item, form, usuario, valores_anteriores=None):
    campos = [
        "tipo",
        "codigo",
        "material",
        "destino",
        "estado_conservacao",
        "quantidade",
        "valor",
        "data_aquisicao",
        "localizacao",
        "status",
    ]

    if valores_anteriores is None:
        valores_anteriores = {
            campo: _normalizar_valor_historico(getattr(item, campo)) for campo in campos
        }

    for campo in campos:
        valor_anterior = valores_anteriores.get(campo, "")
        valor_novo = _normalizar_valor_historico(form.cleaned_data.get(campo))

        if valor_anterior != valor_novo:
            HistoricoAlteracaoPatrimonio.objects.create(
                item=item,
                campo=campo,
                valor_anterior=valor_anterior,
                valor_novo=valor_novo,
                usuario=usuario,
            )


@login_required
def cadptm(request):
    if request.method == "POST":
        form = PatrimonioForm(request.POST)
        if form.is_valid():
            item = form.save()
            HistoricoAlteracaoPatrimonio.objects.create(
                item=item,
                campo="cadastro",
                valor_anterior="",
                valor_novo="item cadastrado",
                usuario=request.user,
            )
            messages.success(request, "PATRIMONIO CADASTRADO COM SUCESSO!")
            return redirect("cadptm")
    else:
        form = PatrimonioForm()

    return render(request, "cadptm.html", {"form": form})


@login_required
def pesquisar(request):
    return render(request, "searchptm.html")


@login_required
def historico_item(request, id):
    item = Patrimonio.objects.filter(id=id).first()
    if not item:
        return HttpResponse(
            "Item não encontrado.", status=404, content_type="text/plain"
        )

    historico = item.historicos.select_related("usuario").all()
    return render(
        request,
        "historico_patrimonio.html",
        {"item": item, "historico": historico},
    )


@login_required
def editar_item(request, id):
    item = Patrimonio.objects.filter(id=id).first()
    if not item:
        return HttpResponse(
            "Item não encontrado.", status=404, content_type="text/plain"
        )
    form = PatrimonioForm(request.POST or None, instance=item)

    if request.method == "POST":
        valores_anteriores = {
            campo: _normalizar_valor_historico(getattr(item, campo))
            for campo in [
                "tipo",
                "codigo",
                "material",
                "destino",
                "estado_conservacao",
                "quantidade",
                "valor",
                "data_aquisicao",
                "localizacao",
                "status",
            ]
        }

        if form.is_valid():
            _registrar_historico_alteracao(item, form, request.user, valores_anteriores)
            form.save()
            messages.success(request, "PATRIMÔNIO ATUALIZADO COM SUCESSO!")
            return redirect("pesquisar")

    return render(request, "cadptm.html", {"form": form, "editar": True, "item": item})


@login_required
def exportar_filtro(request, formato, busca, status):
    itens = Patrimonio.objects.all()

    if busca != "todos":
        itens = itens.filter(
            Q(patrimonio__icontains=busca) | Q(material__icontains=busca)
        )

    if status != "todos":
        itens = itens.filter(status=status)

    ids = ",".join(str(item.id) for item in itens)
    if not ids:
        return HttpResponse(
            "Nenhum item encontrado para exportação.",
            status=404,
            content_type="text/plain",
        )

    return redirect("exportar", formato=formato, ids=ids)


@login_required
def listar_itens(request, busca, status):
    itens = Patrimonio.objects.all()

    material = request.GET.get("material", "todos")
    localizacao = request.GET.get("localizacao", "todos")
    tipo = request.GET.get("tipo", "todos")
    pagina = request.GET.get("pagina", 1)

    if busca != "todos":
        itens = itens.filter(
            Q(patrimonio__icontains=busca) | Q(material__icontains=busca)
        )

    if status != "todos":
        itens = itens.filter(status=status)

    if material != "todos":
        itens = itens.filter(material__icontains=material)

    if localizacao != "todos":
        itens = itens.filter(localizacao__icontains=localizacao)

    if tipo != "todos":
        itens = itens.filter(tipo=tipo)

    paginator = Paginator(itens, 10)
    page_obj = paginator.get_page(pagina)

    return JsonResponse(
        {
            "itens": list(page_obj.object_list.values()),
            "pagina": page_obj.number,
            "total_paginas": paginator.num_pages,
        }
    )


@login_required
@require_POST
def excluir_item(request, id):
    try:
        item = Patrimonio.objects.get(id=id)
        item.delete()

        Auditoria.objects.create(
            usuario=request.user.username,
            acao="EXCLUIR",
            tabela="Patrimonio",
            registro_id=id,
            data_hora=timezone.now(),
        )
        return JsonResponse({"status": "ok"})

    except Patrimonio.DoesNotExist:
        return JsonResponse({"status": "error"}, status=404)
